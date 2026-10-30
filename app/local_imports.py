"""Validated recipient imports and separate SMTP-account CSV imports."""
import csv
from datetime import date, datetime
import io
import re
import zipfile
from openpyxl import load_workbook
from sqlalchemy import func, select
from app.models import Inbox, Lead
from app.schemas import InboxCreate
from app.routers.smtp import SmtpAccountUpsert, upsert_smtp_account
from app.routers.inbox import create_inbox

MAX_BYTES = 20 * 1024 * 1024
MAX_ROWS = 100000
ALIASES = {'почта': 'email', 'электронная почта': 'email', 'имя': 'name', 'компания': 'company',
           'город': 'city', 'оквэд': 'okved', 'должность': 'title', 'источник': 'source',
           'источник контакта': 'source', 'дата согласия': 'consent_date', 'основание согласия': 'consent_basis',
           'сегмент': 'segment', 'лимит': 'max_emails_per_day', 'провайдер': 'provider'}


def email_address(value):
    value = str(value or '').strip().lower()
    if len(value) > 255 or value.count('@') != 1:
        raise ValueError('Некорректный email')
    local, domain = value.rsplit('@', 1)
    try:
        domain = domain.encode('idna').decode('ascii')
    except UnicodeError as exc:
        raise ValueError('Некорректный домен email') from exc
    if not re.fullmatch(r"[a-z0-9!#$%&'*+/=?^_`{|}~.-]{1,64}", local) or local.startswith('.') or local.endswith('.') or '..' in local:
        raise ValueError('Некорректная часть адреса до @')
    if '.' not in domain or any(not re.fullmatch(r'[a-z0-9](?:[a-z0-9-]{0,61}[a-z0-9])?', part) for part in domain.split('.')):
        raise ValueError('Некорректный домен email')
    return local + '@' + domain


def _text(value):
    if value is None:
        return ''
    if isinstance(value, (datetime, date)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    return str(value).strip()


def parse_rows(raw, filename, *, xlsx=True):
    if len(raw) > MAX_BYTES:
        raise ValueError('Максимальный размер файла — 20 МБ')
    if xlsx and filename.lower().endswith('.xlsx'):
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            if len(archive.infolist()) > 2000 or sum(item.file_size for item in archive.infolist()) > 100 * 1024 * 1024:
                raise ValueError('Слишком большой распакованный XLSX')
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True, keep_links=False)
        try:
            iterator = workbook.active.iter_rows(values_only=True)
            header = next(iterator, ())
            values = []
            for row in iterator:
                if len(values) >= MAX_ROWS:
                    raise ValueError('Максимум 100 000 строк за один импорт')
                values.append(row)
        finally:
            workbook.close()
    else:
        try:
            text = raw.decode('utf-8-sig')
        except UnicodeDecodeError:
            text = raw.decode('cp1251')
        try:
            dialect = csv.Sniffer().sniff(text[:8192], delimiters=',;\t')
        except csv.Error:
            dialect = csv.excel
        reader = csv.reader(io.StringIO(text), dialect)
        header = next(reader, ())
        values = []
        for row in reader:
            if len(values) >= MAX_ROWS:
                raise ValueError('Максимум 100 000 строк за один импорт')
            values.append(row)
    keys = [ALIASES.get(_text(key).lower(), _text(key).lower().replace(' ', '_')) for key in header]
    if 'email' not in keys:
        raise ValueError('Нужна колонка email (или Почта)')
    if len(keys) != len(set(keys)):
        raise ValueError('В файле повторяются названия колонок')
    return [{key: _text(row[index]) if index < len(row) else '' for index, key in enumerate(keys) if key}
            for row in values if any(_text(v) for v in row)]


async def import_contacts(db, rows):
    secret_columns = {'smtp_host', 'imap_host', 'smtp_password', 'imap_password', 'password', 'app_password', 'access_token', 'refresh_token', 'api_key'}
    if any(secret_columns.intersection(row) for row in rows):
        raise ValueError('Файл похож на конфигурацию отправителей или содержит секреты. Используйте раздел «Почтовые аккаунты».')
    seen = set(); valid = []; errors = []; duplicates = 0
    for number, row in enumerate(rows, 2):
        try:
            email = email_address(row.get('email'))
            if email in seen:
                duplicates += 1
                continue
            seen.add(email)
            custom = {key: value for key, value in row.items() if key not in {'email', 'name'} and value != ''}
            valid.append((email, row.get('name', ''), custom))
        except ValueError as exc:
            errors.append({'row': number, 'error': str(exc)})
    existing = {}
    for offset in range(0, len(valid), 2000):
        emails = [row[0] for row in valid[offset:offset + 2000]]
        for lead in (await db.execute(select(Lead).where(func.lower(Lead.email).in_(emails)))).scalars().all():
            existing.setdefault(lead.email.lower(), lead)
    added = 0; updated = 0
    for email, name, custom in valid:
        if email in existing:
            lead = existing[email]
            lead.custom_data = {**(lead.custom_data or {}), **custom}
            if name:
                lead.name = name
            updated += 1
        else:
            db.add(Lead(email=email, name=name, custom_data=custom, status='active'))
            added += 1
    await db.flush()
    return {'added': added, 'updated': updated, 'duplicates': duplicates, 'invalid': len(errors), 'errors': errors[:100]}


def _bool(row, key, default):
    value = row.get(key, '')
    return default if value == '' else value.strip().lower() in {'1', 'true', 'yes', 'да'}


async def import_smtp_accounts(db, rows):
    from app.local_send import local_install
    added = 0; errors = []; seen = set()
    for number, row in enumerate(rows, 2):
        try:
            email = email_address(row.get('email'))
            if email in seen:
                raise ValueError('Повторяющийся email в файле')
            seen.add(email)
            if row.get('provider', 'smtp').lower() in {'gmail', 'google', 'office365', 'outlook', 'microsoft'} or email.rsplit('@', 1)[1] in {'gmail.com', 'googlemail.com', 'outlook.com', 'hotmail.com', 'live.com'}:
                raise ValueError('Google и Microsoft подключаются через OAuth в интерфейсе')
            limit = int(row.get('max_emails_per_day') or 5)
            if not 1 <= limit <= 100000:
                raise ValueError('Дневной лимит должен быть от 1 до 100000 и соответствовать правилам вашего провайдера')
            async with db.begin_nested():
                inbox = await db.scalar(select(Inbox).where(func.lower(Inbox.email) == email))
                if inbox is None:
                    inbox = await create_inbox(InboxCreate(email=email, provider='smtp', max_emails_per_day=limit,
                                               display_name=row.get('display_name', '')), db=db)
                elif inbox.provider != 'smtp':
                    raise ValueError('Этот адрес уже подключён через OAuth')
                inbox.max_emails_per_day = limit
                account = SmtpAccountUpsert(smtp_host=row.get('smtp_host', ''), smtp_port=int(row.get('smtp_port') or 465),
                    smtp_username=row.get('smtp_username') or email, smtp_password=row.get('smtp_password', ''),
                    smtp_use_ssl=_bool(row, 'smtp_use_ssl', int(row.get('smtp_port') or 465) == 465),
                    smtp_use_tls=_bool(row, 'smtp_use_tls', int(row.get('smtp_port') or 465) != 465),
                    imap_host=row.get('imap_host', ''), imap_port=int(row.get('imap_port') or 993),
                    imap_username=row.get('imap_username') or email, imap_password=row.get('imap_password', ''),
                    imap_use_ssl=_bool(row, 'imap_use_ssl', True))
                await upsert_smtp_account(inbox.id, account, db=db, _user=object())
            added += 1
        except Exception as exc:
            # Never return Pydantic reprs or submitted values: they may contain passwords.
            from fastapi import HTTPException
            from pydantic import ValidationError
            if isinstance(exc, HTTPException):
                detail = str(exc.detail)
            elif isinstance(exc, ValidationError):
                detail = 'Некорректные поля SMTP/IMAP (проверьте порты и обязательные значения)'
            elif isinstance(exc, ValueError):
                detail = 'Некорректные числовые поля' if 'invalid literal' in str(exc) else str(exc)
            else:
                detail = 'Не удалось сохранить строку'
            errors.append({'row': number, 'error': detail[:1000]})
    await db.flush()
    return {'saved': added, 'invalid': len(errors), 'errors': errors[:100]}
