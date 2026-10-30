"""Statistics, license-aware field projections and spreadsheet-safe downloads."""
import csv
import io
from sqlalchemy import func, select
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from app.collector.models import Company, CompanyContact, CompanyObservation, CollectorSource, CollectionJob, CollectionResult
from app.collector.repository import FIELDS


def safe_cell(value):
    value = str(value if value is not None else '')
    if value.lstrip().startswith(('=', '+', '-', '@', '\t', '\r')):
        return "'" + value
    return value


async def exportable_company(db, company):
    observations = list((await db.scalars(select(CompanyObservation).where(CompanyObservation.company_id == company.id).order_by(CompanyObservation.checked_at))).all())
    sources = {s.key: s for s in (await db.scalars(select(CollectorSource))).all()}
    def allowed(origin):
        source = sources.get(origin.get('source'))
        rights = origin.get('rights', {})
        return source and source.allow_store and source.allow_export and rights.get('store') and rights.get('export')
    projection, urls, provenance = {}, [], []
    for observation in observations:
        if not allowed({'source': observation.source, 'rights': observation.rights}): continue
        for key in FIELDS:
            value = (observation.fields or {}).get(key)
            if value not in ('', None, []): projection[key] = value
        if observation.url: urls.append(observation.url)
        provenance.append(observation.source)
    if not projection.get('name'):
        return None
    projection['emails'], projection['phones'] = [], []
    for contact in (await db.scalars(select(CompanyContact).where(CompanyContact.company_id == company.id))).all():
        origins = [o for o in contact.origins if allowed(o)]
        if not origins: continue
        if contact.kind in {'email', 'phone'}:
            projection['emails' if contact.kind == 'email' else 'phones'].append(contact.value)
            urls.extend(o['url'] for o in origins if o.get('url'))
    projection.update(id=company.id, status=company.status, sources=sorted(set(provenance)),
                      source_urls=sorted(set(urls)), checked_at=company.checked_at.isoformat())
    return projection


async def statistics(db, company_ids=None):
    company_query = select(Company.id)
    if company_ids is not None: company_query = company_query.where(Company.id.in_(company_ids))
    ids = company_query
    total = await db.scalar(select(func.count()).select_from(Company).where(Company.id.in_(ids))) or 0
    emails = await db.scalar(select(func.count(func.distinct(CompanyContact.value))).where(CompanyContact.company_id.in_(ids), CompanyContact.kind == 'email')) or 0
    phones = await db.scalar(select(func.count(func.distinct(CompanyContact.value))).where(CompanyContact.company_id.in_(ids), CompanyContact.kind == 'phone')) or 0
    with_email = await db.scalar(select(func.count(func.distinct(CompanyContact.company_id))).where(CompanyContact.company_id.in_(ids), CompanyContact.kind == 'email')) or 0
    with_contacts = await db.scalar(select(func.count(func.distinct(CompanyContact.company_id))).where(CompanyContact.company_id.in_(ids))) or 0
    return {'companies': total, 'unique_emails': emails, 'phones': phones, 'with_email': with_email,
            'without_email': total - with_email, 'without_contacts': total - with_contacts}


EXPORT_FIELDS = ['emails', 'name', 'phones', 'website', 'city', 'address',
    *[field for field in FIELDS if field not in {'name', 'website', 'city', 'address'}],
    'source_urls', 'sources', 'checked_at', 'status', 'id']
EXPORT_LABELS = {'emails': 'Email', 'name': 'Компания', 'phones': 'Телефоны', 'website': 'Сайт',
    'city': 'Город', 'address': 'Адрес', 'inn': 'ИНН', 'kpp': 'КПП', 'ogrn': 'ОГРН',
    'okved_main': 'ОКВЭД', 'okved_extra': 'Другие ОКВЭД', 'okved_as_of': 'Дата ОКВЭД',
    'country': 'Страна', 'region': 'Регион', 'industry': 'Деятельность', 'employees': 'Численность',
    'active': 'Действующая', 'source_urls': 'Страницы источников', 'sources': 'Источники',
    'checked_at': 'Дата проверки', 'status': 'Статус', 'id': 'ID'}


async def export_bytes(db, companies, format):
    rows = []
    omitted = 0
    for company in companies:
        row = await exportable_company(db, company)
        if not row or not row['emails']:
            omitted += 1
            continue
        rows.append([safe_cell('; '.join(map(str, row.get(field, []))) if isinstance(row.get(field), list) else row.get(field, '')) for field in EXPORT_FIELDS])
    if format == 'csv':
        stream = io.StringIO(newline='')
        writer = csv.writer(stream)
        writer.writerow(EXPORT_FIELDS)
        writer.writerows(rows)
        return stream.getvalue().encode('utf-8-sig'), omitted
    book = Workbook(write_only=True)
    sheet = book.create_sheet('Компании')
    sheet.freeze_panes = 'C2'
    sheet.auto_filter.ref = f'A1:{get_column_letter(len(EXPORT_FIELDS))}{len(rows) + 1}'
    for column, width in [('A', 48), ('B', 40), ('C', 24), ('D', 36), ('E', 22), ('F', 50)]:
        sheet.column_dimensions[column].width = width
    sheet.append([EXPORT_LABELS[field] for field in EXPORT_FIELDS])
    for row in rows: sheet.append(row)
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue(), omitted
