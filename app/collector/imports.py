"""CSV/XLSX (including GeoLeadScraper) and bounded official FNS XML import."""
import csv
import io
import json
import zipfile
from datetime import date, datetime
import xml.etree.ElementTree as ET
from openpyxl import load_workbook
from app.collector.validation import split_values

MAX_BYTES = 20 * 1024 * 1024
MAX_ROWS = 10000
ALIASES = {'title': 'name', 'название': 'name', 'наименование': 'name', 'компания': 'name',
    'инн': 'inn', 'кпп': 'kpp', 'огрн': 'ogrn', 'регион': 'region', 'город': 'city', 'municipality': 'city',
    'адрес': 'address', 'сайт': 'website', 'email': 'email', 'почта': 'email', 'телефон': 'phone',
    'оквэд': 'okved_main', 'основной_оквэд': 'okved_main', 'дополнительные_оквэд': 'okved_extra',
    'categories': 'industry', 'maps_url': 'source_url', 'place_id': 'external_id',
    'сотрудников': 'employees', 'действующая': 'active', 'источник': 'source', 'url_источника': 'source_url'}
ALLOWED = {'name', 'inn', 'kpp', 'ogrn', 'region', 'city', 'country', 'address', 'website', 'email', 'emails',
    'phone', 'phones', 'industry', 'okved_main', 'okved_extra', 'okved_as_of', 'source_url', 'source',
    'external_id', 'active', 'employees'}


def text(value):
    if value is None: return ''
    if isinstance(value, (date, datetime)): return value.isoformat()
    return str(value).strip()


def parse_company_file(raw, filename):
    if len(raw) > MAX_BYTES:
        raise ValueError('Предел файла — 20 МБ')
    if filename.lower().endswith('.xml'):
        return parse_fns_xml(raw)
    if filename.lower().endswith('.xlsx'):
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            if len(archive.infolist()) > 2000 or sum(i.file_size for i in archive.infolist()) > 100 * 1024 * 1024:
                raise ValueError('Слишком большой распакованный XLSX')
        book = load_workbook(io.BytesIO(raw), read_only=True, data_only=True, keep_links=False)
        try:
            rows = book.active.iter_rows(values_only=True)
            header = next(rows, [])
            values = []
            for row in rows:
                values.append(row)
                if len(values) > MAX_ROWS: raise ValueError('Не более 10 000 строк за импорт')
        finally:
            book.close()
    elif filename.lower().endswith(('.csv', '.tsv')):
        try: data = raw.decode('utf-8-sig')
        except UnicodeDecodeError: data = raw.decode('cp1251')
        try: dialect = csv.Sniffer().sniff(data[:8192], delimiters=',;\t')
        except csv.Error: dialect = csv.excel
        reader = csv.reader(io.StringIO(data), dialect)
        header, values = next(reader, []), []
        for row in reader:
            values.append(row)
            if len(values) > MAX_ROWS: raise ValueError('Не более 10 000 строк за импорт')
    else:
        raise ValueError('Поддерживаются CSV, TSV, XLSX и XML ЕГРЮЛ/ЕГРИП')
    keys = [ALIASES.get(text(v).lower().replace(' ', '_'), text(v).lower().replace(' ', '_')) for v in header]
    if 'name' not in keys:
        raise ValueError('Нужна колонка name, title или Название')
    if len(set(keys)) != len(keys): raise ValueError('Повторяющиеся названия колонок')
    if any(any(secret in k for secret in ['password', 'token', 'api_key', 'пароль']) for k in keys):
        raise ValueError('Файл компаний не должен содержать пароли, ключи и токены')
    records = []
    for values_row in values:
        if not any(text(v) for v in values_row): continue
        record = {k: text(values_row[i]) if i < len(values_row) else '' for i, k in enumerate(keys) if k in ALLOWED}
        for key in ['emails', 'phones', 'okved_extra']:
            value = record.get(key, '')
            if value.startswith('['):
                try: record[key] = json.loads(value)
                except ValueError: pass
        url = record.get('source_url', '')
        # Keep the original provider attribution for licensed map exports.
        if 'yandex.' in url: record['source'] = 'yandex'
        elif '2gis.' in url: record['source'] = 'twogis'
        records.append(record)
    return records


def parse_fns_xml(raw):
    if b'<!DOCTYPE' in raw.upper() or b'<!ENTITY' in raw.upper() or b'\x00' in raw[:100]:
        raise ValueError('DTD, внешние сущности и UTF-16 XML не поддерживаются')
    records = []
    for _, element in ET.iterparse(io.BytesIO(raw), events=['end']):
        tag = element.tag.rsplit('}', 1)[-1]
        if tag not in {'СвЮЛ', 'СвИП'}: continue
        attrs = element.attrib
        def nodes(name):
            return [e for e in element.iter() if e.tag.rsplit('}', 1)[-1] == name]
        names = nodes('СвНаимЮЛ')
        name = names[0].get('НаимЮЛПолн', '') if names else ''
        if not name and tag == 'СвИП':
            # Do not import private names from registry files as company contacts.
            name = 'ИП ' + attrs.get('ОГРНИП', '')
        main = nodes('СвОКВЭДОсн')
        row = {'source': 'fns', 'name': name, 'inn': attrs.get('ИНН', attrs.get('ИННФЛ', '')),
            'ogrn': attrs.get('ОГРН', attrs.get('ОГРНИП', '')), 'kpp': attrs.get('КПП', ''),
            'okved_main': main[0].get('КодОКВЭД', '') if main else '',
            'okved_extra': [e.get('КодОКВЭД', '') for e in nodes('СвОКВЭДДоп')],
            'okved_as_of': attrs.get('ДатаВып', ''), 'active': not bool(nodes('СвПрекрЮЛ') or nodes('СвПрекращ')),
            'external_id': attrs.get('ОГРН', attrs.get('ОГРНИП', '')),
            'source_url': 'https://egrul.nalog.ru/'}
        address_nodes = nodes('АдресРФ') or nodes('СвАдрЮЛФИАС')
        if address_nodes:
            a = address_nodes[0]
            row['address'] = ', '.join(v for e in a.iter() for k, v in e.attrib.items() if k in {'Индекс', 'НаимРегион', 'НаимГород', 'НаимНаселПункт', 'НаимУлица', 'Дом', 'Корпус', 'Кварт'})
            for e in a.iter():
                if e.tag.endswith('Регион'): row['region'] = e.get('НаимРегион', '')
                if e.tag.endswith('Город'): row['city'] = e.get('НаимГород', '')
        records.append(row)
        element.clear()
        if len(records) > MAX_ROWS: raise ValueError('Разделите XML на части до 10 000 организаций')
    if not records: raise ValueError('Не найдены записи СвЮЛ или СвИП в XML ФНС')
    return records
