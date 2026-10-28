"""Optional import of the services' official CSV/XLSX exports."""
from __future__ import annotations

import csv
import hashlib
import io
import json
import re
from pathlib import Path

from .sources import source_url
from .storage import Store, contact_url
from .websites import extract_emails

ALIASES = {
    "name": {"наименование", "название", "компания", "организация", "наименованиеорганизации", "наименованиекомпании", "полноенаименование", "краткоенаименование", "полноенаименованиеорганизации", "наименованиеюл", "наименованиеорганизацииип", "наименованиеконтрагента", "name", "company"},
    "inn": {"инн", "inn"}, "address": {"адрес", "юридическийадрес", "адресместонахождения", "address"},
    "city": {"город", "населенныйпункт", "населённыйпункт", "city"}, "region": {"регион", "субъект", "region"},
    "rubrics": {"оквэд", "основнойоквэд", "основнойвиддеятельности", "виддеятельности", "rubrics"},
    "source_url": {"ссылка", "ссылканакарточку", "карточкаисточника", "sourceurl"},
}


def heading(value) -> str:
    return re.sub(r"[^a-zа-яё0-9]", "", str(value or "").casefold())


def field(header: str) -> str | None:
    key = heading(header)
    for kind, names in ALIASES.items():
        if key in names:
            return kind
    if re.match(r"(?:email|электроннаяпочта|адресэлектроннойпочты|электронныйадрес|почта)\d*$", key):
        return "emails"
    if re.match(r"(?:сайт|сайты|вебсайт|website|websites)\d*$", key):
        return "websites"
    if re.match(r"(?:телефон|телефоны|phone)\d*$", key):
        return "phones"
    return None


def table_rows(path: Path):
    if path.suffix.lower() == ".csv":
        raw = path.read_bytes()
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = raw.decode("cp1251")
        try:
            dialect = csv.Sniffer().sniff(text[:10000], delimiters=";,\t")
        except csv.Error:
            dialect = csv.excel
        yield from csv.reader(io.StringIO(text), dialect)
    elif path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            yield from workbook.worksheets[0].iter_rows(values_only=True)
        finally:
            workbook.close()
    else:
        raise ValueError("Поддерживаются CSV и XLSX. Старый XLS сначала сохраните как XLSX или CSV.")


def records(path: Path, source: str, query: str):
    mapping = None
    for number, row in enumerate(table_rows(path), 1):
        if mapping is None:
            kinds = [field(cell) for cell in row]
            if "name" in kinds:
                mapping = kinds
            elif number >= 50:
                raise ValueError("Не найдена строка заголовков с названием компании. Нужна таблица компаний с колонками ИНН/Email/Сайт.")
            continue
        values = {key: [] for key in (*ALIASES, "emails", "websites", "phones")}
        for kind, value in zip(mapping, row):
            if kind and value is not None:
                values[kind].append(str(value).strip())
        if not values["name"] or not values["name"][0]:
            continue
        name = values["name"][0]
        inn = values["inn"][0] if values["inn"] else ""
        inn = inn if re.fullmatch(r"\d{10}|\d{12}", inn) else ""
        address = " | ".join(values["address"])
        websites = set()
        for value in values["websites"]:
            for candidate in re.split(r"[\s;,|]+", value):
                if candidate and (url := contact_url(candidate)):
                    websites.add(url)
        link = ""
        if values["source_url"]:
            try:
                link = source_url(source, values["source_url"][0])
            except ValueError:
                pass
        yield {"schema": "company-contact/v1", "source": source,
               "external_id": inn or hashlib.sha256((name + "|" + address).encode()).hexdigest()[:24],
               "name": name, "inn": inn, "address": address,
               "city": " | ".join(values["city"]), "region": " | ".join(values["region"]),
               "rubrics": values["rubrics"] or [query], "emails": sorted({e for v in values["emails"] for e in extract_emails(v) if "*" not in e}),
               "websites": sorted(websites), "phones": values["phones"], "source_url": link, "origin_row": number}
    if mapping is None:
        raise ValueError("Файл пуст или не содержит заголовка компаний")


def import_file(store: Store, source: str, path: Path, query: str) -> int:
    path = path.expanduser().resolve()
    if not path.is_file():
        raise ValueError("Файл выгрузки не найден")
    identity = hashlib.sha256(source.encode() + path.read_bytes()).hexdigest()[:20]
    task = {"id": identity, "source": source, "query": query, "city": "Россия",
            "import_file": str(path), "url": str(path), "limit": None}
    store.add_tasks([task])
    folder = store.directory / "jobs" / identity
    folder.mkdir(parents=True, exist_ok=True)
    raw = folder / "import.jsonl"
    store.mark(identity, "running")
    try:
        with raw.open("w", encoding="utf-8") as output:
            for record in records(path, source, query):
                output.write(json.dumps(record, ensure_ascii=False) + "\n")
        count = store.import_items(task, raw)
        if not count:
            raise ValueError("В выгрузке не найдены компании")
        store.mark(identity, "done", count)
        print(f"Импортировано компаний: {count}")
        return count
    except BaseException:
        store.mark(identity, "error", message="Импорт не завершён; проверьте формат таблицы")
        raise
