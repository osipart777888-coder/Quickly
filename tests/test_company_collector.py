"""Collector safety, persistence and integration tests. No external delivery/API IO."""
import asyncio
import csv
import io
import time
from datetime import timedelta
from types import SimpleNamespace
import pytest
from sqlalchemy import select, func, cast, Text
from openpyxl import Workbook, load_workbook
from app.collector.models import Company, CompanyContact, CompanyObservation, CollectionJob, CollectionTask, ContactPermission, CollectorSource
from app.collector.repository import save_company
from app.collector.sources import ensure_sources, permissions, TwoGisAdapter, reserve_api_request
from app.collector.websites import Page, SourceError, extract_website, public_addresses
from app.collector.imports import parse_company_file
from app.collector.consent import set_permission, transfer_to_quickly, marketing_error
from app.collector.reports import export_bytes
from app.collector import queue
from app.collector.validation import valid_inn, contact_email
from app.models import Lead, Campaign, CampaignLead
from app.time import utcnow

RIGHTS = {'store': True, 'export': True, 'basis': 'Test fixture owned by Quickly', 'terms_url': 'https://fixture.example/terms'}


def record(index=1, **overrides):
    return {'name': f'Компания {index}', 'website': f'https://company-{index}.test-fixture.ru/',
            'email': f'office{index}@test-fixture.ru', 'city': 'Москва', 'active': True, **overrides}


@pytest.mark.asyncio
async def test_identity_provenance_and_distinct_legal_entities(session):
    await ensure_sources(session)
    a, _ = await save_company(session, record(1, inn='7707083893', kpp='770701001'), source='import', rights=RIGHTS)
    again, duplicate = await save_company(session, record(1, inn='7707083893', kpp='770701001'), source='import', rights=RIGHTS)
    assert duplicate and a.id == again.id
    b, _ = await save_company(session, record(1, inn='7736050003', kpp='773601001'), source='import', rights=RIGHTS)
    assert b.id != a.id
    branch, _ = await save_company(session, record(1, inn='7707083893', kpp='770702001'), source='fns', rights=RIGHTS)
    assert branch.id not in {a.id, b.id}
    assert await session.scalar(select(func.count()).select_from(Company)) == 3
    assert await session.scalar(select(func.count()).select_from(CompanyObservation)) == 3
    assert (await session.get(ContactPermission, 'office1@test-fixture.ru')).status == 'unknown'


@pytest.mark.asyncio
async def test_license_blocks_storage_and_revocation_blocks_export(session):
    await ensure_sources(session)
    source = await session.get(CollectorSource, 'yandex')
    with pytest.raises(SourceError): permissions(source)
    with pytest.raises(ValueError): await save_company(session, record(), source='yandex', rights={'store': False})
    company, _ = await save_company(session, record(), source='import', rights=RIGHTS)
    raw, omitted = await export_bytes(session, [company], 'csv')
    assert omitted == 0 and b'office1@test-fixture.ru' in raw
    source = await session.get(CollectorSource, 'import')
    source.allow_export = False
    await session.flush()
    raw, omitted = await export_bytes(session, [company], 'csv')
    assert omitted == 1 and b'office1@test-fixture.ru' not in raw


def test_website_robots_contact_pages_and_no_external_crawls():
    calls = []
    pages = {
        '/robots.txt': 'User-agent: *\nDisallow: /private\nCrawl-delay: 2',
        '/': '<a href="/contacts">Контакты</a><a href="https://other.invalid/about">О нас</a><a href="/private/contact">Контакты</a><script>secret@do-not-extract.ru</script>',
        '/contacts': '<a href="mailto:Sales@Test-Fixture.ru?subject=hello">email</a> office@test-fixture.ru <a href="tel:+7(495)123-45-67">Телефон</a><form action="/send"><input name="email"></form>',
    }
    def fetch(url, **kwargs):
        from urllib.parse import urlsplit
        path = urlsplit(url).path
        calls.append(path)
        return Page(url, 200, {'content-type': 'text/html'}, pages[path])
    result = extract_website('https://test-fixture.ru/', fetch=fetch, sleep=lambda seconds: None)
    assert calls == ['/robots.txt', '/', '/contacts']
    assert {'sales@test-fixture.ru', 'office@test-fixture.ru'} <= {c['value'] for c in result['contacts']}
    assert all(c['url'].endswith('/contacts') for c in result['contacts'])
    assert any(c['kind'] == 'form' for c in result['contacts'])
    assert '+74951234567' in {c['value'] for c in result['contacts']}


def test_robots_disallow_and_captcha_stop():
    def deny(url, **kwargs): return Page(url, 200, {}, 'User-agent: *\nDisallow: /')
    with pytest.raises(SourceError, match='robots'): extract_website('https://test-fixture.ru/', fetch=deny, sleep=lambda _: None)
    def captcha(url, **kwargs):
        return Page(url, 200, {'content-type': 'text/html'}, '' if url.endswith('robots.txt') else '<title>CAPTCHA</title>')
    with pytest.raises(SourceError, match='CAPTCHA'): extract_website('https://test-fixture.ru/', fetch=captcha, sleep=lambda _: None)


def test_ssrf_blocks_private_and_mixed_dns(monkeypatch):
    import socket
    monkeypatch.setattr(socket, 'getaddrinfo', lambda *args, **kwargs: [(2, 1, 6, '', ('8.8.8.8', 0)), (2, 1, 6, '', ('127.0.0.1', 0))])
    with pytest.raises(SourceError): public_addresses('test-fixture.ru')
    with pytest.raises(SourceError): public_addresses('localhost')


def test_vpn_synthetic_dns_fallback_still_checks_public_destination(monkeypatch):
    import socket
    monkeypatch.setattr(socket, 'getaddrinfo', lambda *a, **k: [(2,1,6,'',('198.18.0.1',0))])
    monkeypatch.setattr('app.collector.websites.doh_addresses', lambda host: ['8.8.8.8'])
    assert public_addresses('fixture.ru') == ['8.8.8.8']
    monkeypatch.setattr('app.collector.websites.doh_addresses', lambda host: ['127.0.0.1'])
    with pytest.raises(SourceError): public_addresses('fixture.ru')


def test_csv_xlsx_geolead_compatibility_and_formula_protection():
    data = b'title,website,email,phone,maps_url,place_id,categories\nCompany,https://company.ru,office@company.ru,+74951234567,https://2gis.ru/firm/123,123,Software\n'
    rows = parse_company_file(data, 'geolead.csv')
    assert rows[0]['source'] == 'twogis'
    assert rows[0]['external_id'] == '123'
    assert rows[0]['industry'] == 'Software'
    assert not rows[0].get('okved_main')
    book = Workbook(); sheet = book.active
    sheet.append(['Название', 'Сайт', 'Почта']); sheet.append(['ООО Пример', 'https://company.ru', 'office@company.ru'])
    stream = io.BytesIO(); book.save(stream)
    assert parse_company_file(stream.getvalue(), 'companies.xlsx')[0]['name'] == 'ООО Пример'
    with pytest.raises(ValueError): parse_company_file(b'name,password\nx,secret', 'bad.csv')


def test_fns_exact_okved_and_xxe_rejected():
    xml = '''<?xml version="1.0" encoding="UTF-8"?><Файл><СвЮЛ ИНН="7707083893" КПП="770701001" ОГРН="1027700132195" ДатаВып="2026-10-26"><СвНаимЮЛ НаимЮЛПолн="Тест"/><СвОКВЭД><СвОКВЭДОсн КодОКВЭД="62.01"/><СвОКВЭДДоп КодОКВЭД="62.02"/></СвОКВЭД></СвЮЛ></Файл>'''.encode()
    row = parse_company_file(xml, 'fns.xml')[0]
    assert row['okved_main'] == '62.01' and row['okved_extra'] == ['62.02']
    assert row['source'] == 'fns' and row['active'] is True
    with pytest.raises(ValueError): parse_company_file(b'<!DOCTYPE x [<!ENTITY x SYSTEM "file:///etc/passwd">]><x>&x;</x>', 'bad.xml')


@pytest.mark.asyncio
async def test_paid_budget_debit_and_no_contact_field_without_explicit_access(session, monkeypatch):
    await ensure_sources(session)
    source = await session.get(CollectorSource, 'twogis')
    source.enabled = source.allow_store = source.allow_export = source.paid_authorized = True
    source.rights_basis = 'Fixture contract'; source.api_key = 'fixture-key-not-live'
    source.daily_request_limit = 1
    source.config = {'region_id': '123', 'daily_budget_rub': 1, 'request_cost_rub': 1}
    await session.commit()
    await reserve_api_request(session, 'twogis')
    source.next_request_at = None
    await session.commit()
    with pytest.raises(SourceError, match='бюджет'): await reserve_api_request(session, 'twogis')
    captured = {}
    def get_json(url, params):
        captured.update(params)
        return {'meta': {'code': 200}, 'result': {'items': []}}
    monkeypatch.setattr('app.collector.sources.get_json', get_json)
    TwoGisAdapter().search(source, 'компании Москва')
    assert 'items.contact_groups' not in captured['fields']
    source.config = {**source.config, 'contacts_access': True}
    TwoGisAdapter().search(source, 'компании Москва')
    assert 'items.contact_groups' in captured['fields']


@pytest.mark.asyncio
async def test_consent_required_transfer_paused_and_optout_blocks_future(session):
    await ensure_sources(session)
    company, _ = await save_company(session, record(), source='import', rights=RIGHTS)
    email = 'office1@test-fixture.ru'
    result = await transfer_to_quickly(session, [company.id])
    assert result['created'] == 0 and result['blocked'] == 1
    assert await marketing_error(session, email)
    with pytest.raises(ValueError): await set_permission(session, email, status='confirmed')
    await set_permission(session, email, status='confirmed', advertiser='Syntolk fixture', evidence='Signed consent fixture', evidence_url='consent-001', granted_at=utcnow()-timedelta(days=1))
    campaign = Campaign(name='Fixture campaign'); session.add(campaign); await session.flush()
    result = await transfer_to_quickly(session, [company.id], campaign_id=campaign.id)
    assert result['created'] == 1 and result['enrolled'] == 1
    assert (await session.scalar(select(CampaignLead))).sending_paused is True
    await set_permission(session, email, status='refused', evidence='Fixture opt out')
    assert await marketing_error(session, email)
    assert (await session.scalar(select(Lead))).status == 'unsubscribed'
    assert (await transfer_to_quickly(session, [company.id]))['blocked'] == 1
    with pytest.raises(ValueError, match='отказов'):
        await set_permission(session, email, status='confirmed', advertiser='fixture', evidence='fixture', evidence_url='doc', granted_at=utcnow())


@pytest.mark.asyncio
async def test_export_neutralizes_formulas_and_xlsx_has_all_contacts(session):
    await ensure_sources(session)
    company, _ = await save_company(session, record(name='=CMD()', emails=['sales@test-fixture.ru']), source='import', rights=RIGHTS)
    raw, _ = await export_bytes(session, [company], 'csv')
    assert "'=CMD()" in raw.decode('utf-8-sig')
    raw, _ = await export_bytes(session, [company], 'xlsx')
    book = load_workbook(io.BytesIO(raw)); rows = list(book.active.values)
    assert rows[1][1] == "'=CMD()"
    assert 'office1@test-fixture.ru' in str(rows[1]) and 'sales@test-fixture.ru' in str(rows[1])


@pytest.mark.asyncio
async def test_queue_pause_resume_and_expired_lease_recovery(engine, session):
    from app.routers.collector import job_action
    params = {'sources': ['import'], 'rights_confirmed': True, 'rights_basis': 'Fixture', 'limit': 100, 'crawl_websites': False}
    job = await queue.create_job(session, params, records=[record(i) for i in range(30)])
    await session.commit()
    await job_action(job.id, 'pause', db=session, _admin=None); await session.commit()
    await queue.run_collection_tick(max_tasks=10)
    assert await session.scalar(select(func.count()).select_from(Company)) == 0
    await job_action(job.id, 'resume', db=session, _admin=None); await session.commit()
    task = await queue.claim_task(session)
    assert task and task.status == 'running'
    task.lease_until = utcnow()-timedelta(seconds=1); await session.commit()
    await queue.run_collection_tick(max_tasks=10)
    await session.refresh(job)
    assert job.status == 'completed'
    assert await session.scalar(select(func.count()).select_from(Company)) == 30
    assert all(t.payload == {} for t in (await session.scalars(select(CollectionTask).execution_options(populate_existing=True))).all())


@pytest.mark.asyncio
async def test_one_thousand_synthetic_records_import_queue_and_dedup(engine, session):
    params = {'sources': ['import'], 'rights_confirmed': True, 'rights_basis': 'Synthetic benchmark', 'limit': 1000, 'crawl_websites': False}
    csvfile = io.StringIO(); writer = csv.DictWriter(csvfile, fieldnames=['name', 'website', 'email', 'city', 'active']); writer.writeheader()
    writer.writerows(record(i) for i in range(1000))
    rows = parse_company_file(csvfile.getvalue().encode(), 'synthetic.csv')
    started = time.monotonic()
    job = await queue.create_job(session, params, records=rows); await session.commit()
    await queue.run_collection_tick(max_tasks=100)
    await session.refresh(job)
    assert job.status == 'completed'
    assert await session.scalar(select(func.count()).select_from(Company)) == 1000
    repeat = await queue.create_job(session, params, records=rows); await session.commit()
    await queue.run_collection_tick(max_tasks=100)
    await session.refresh(repeat)
    assert repeat.duplicates == 1000
    assert await session.scalar(select(func.count()).select_from(Company)) == 1000
    assert await session.scalar(select(func.count()).select_from(CompanyContact)) == 1000
    print(f'1000 synthetic records + replay: {time.monotonic()-started:.2f}s')


@pytest.mark.asyncio
async def test_new_api_requires_auth_and_never_returns_keys(session):
    import httpx
    from fastapi import FastAPI
    from app.routers.collector import router
    from app.auth import get_current_user
    from app.database import get_db
    app = FastAPI(); app.include_router(router)
    async def db(): yield session
    app.dependency_overrides[get_db] = db
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url='http://test') as client:
        assert (await client.get('/api/collector/sources')).status_code == 401
        app.dependency_overrides[get_current_user] = lambda: SimpleNamespace(id=1, role='admin')
        assert (await client.get('/api/collector/sources')).status_code == 200
        source = await session.get(CollectorSource, 'yandex'); source.api_key = 'not-a-real-fixture-api-key'; await session.flush()
        response = await client.get('/api/collector/sources')
        assert 'not-a-real-fixture-api-key' not in response.text
        assert 'has_api_key' in response.text
        response = await client.post('/api/collector/jobs', json={'sources':['website'], 'websites':['https://fixture.ru/'], 'rights_confirmed':False})
        assert response.status_code == 400


@pytest.mark.asyncio
async def test_consent_guard_blocks_campaign_before_transport(session, monkeypatch):
    from app.local_send import guarded_send
    from app.sender import SendFailure
    session.add(ContactPermission(email='new@test-fixture.ru', status='unknown'))
    await session.flush()
    calls = []
    def send(**kwargs): calls.append(kwargs)
    monkeypatch.setenv('QUICKLY_LOCAL_INSTALL', 'false')
    result = await guarded_send(session, SimpleNamespace(id=1), kind='campaign', send=send, to_email='new@test-fixture.ru')
    assert isinstance(result, SendFailure)
    assert calls == []


@pytest.mark.asyncio
async def test_retry_backoff_survives_sessions_and_is_bounded(engine, session, monkeypatch):
    params = {'sources':['website'], 'rights_confirmed':True, 'rights_basis':'Fixture', 'limit':1, 'websites':['https://site.test-fixture.ru/']}
    job = await queue.create_job(session, params); await session.commit()
    def unavailable(*args, **kwargs): raise SourceError('Fixture 429', retry_after=60, blocked=True)
    monkeypatch.setattr(queue, 'extract_website', unavailable)
    await queue.run_collection_tick(max_tasks=2)
    task = await session.scalar(select(CollectionTask).where(CollectionTask.kind == 'website').execution_options(populate_existing=True))
    assert task.status == 'retry' and task.due_at > utcnow()
    for _ in range(2):
        task.due_at = utcnow()-timedelta(seconds=1)
        from app.collector.models import WebsiteThrottle
        throttle = await session.get(WebsiteThrottle, 'site.test-fixture.ru')
        throttle.next_request_at = utcnow()-timedelta(seconds=1)
        await session.commit()
        await queue.run_collection_tick(max_tasks=1)
        await session.refresh(task)
    assert task.status == 'error' and task.attempts == 3
    await session.refresh(job)
    assert job.status == 'completed_with_errors'


@pytest.mark.asyncio
async def test_field_level_license_projection_and_erasure(session):
    from app.routers.collector import erase_contact
    await ensure_sources(session)
    company, _ = await save_company(session, record(), source='import', rights=RIGHTS)
    await save_company(session, record(address='SECRET LICENSED ADDRESS', phone='+74959876543', external_id='restricted'), source='yandex', rights={'store':True,'export':False})
    raw, _ = await export_bytes(session, [company], 'csv')
    assert 'SECRET LICENSED ADDRESS' not in raw.decode('utf-8-sig')
    assert b'74959876543' not in raw
    email = 'office1@test-fixture.ru'
    await set_permission(session, email, status='confirmed', advertiser='Fixture', evidence='signed fixture', evidence_url='DOC-1', granted_at=utcnow())
    await transfer_to_quickly(session, [company.id])
    contact = await session.scalar(select(CompanyContact).where(CompanyContact.value == email))
    await erase_contact(contact.id, db=session, _admin=None)
    assert not await session.get(ContactPermission, email, populate_existing=True)
    assert not await session.scalar(select(Lead.id).where(Lead.email == email))
    assert not await session.scalar(select(CompanyContact.id).where(CompanyContact.value == email))
    await save_company(session, record(), source='import', rights=RIGHTS)
    assert not await session.scalar(select(CompanyContact.id).where(CompanyContact.value == email))


@pytest.mark.asyncio
async def test_source_secrets_encrypted_at_rest(session, monkeypatch):
    from app import security
    old = security._fernet
    security.init_encryption('collector-test-fixture-key')
    try:
        source = CollectorSource(key='fixture', api_key='test-secret-must-be-encrypted')
        session.add(source); await session.flush()
        stored = await session.scalar(select(cast(CollectorSource.api_key, Text)).where(CollectorSource.key == 'fixture'))
        assert stored.startswith('gAAAAA') and 'test-secret' not in stored
    finally:
        security._fernet = old


@pytest.mark.asyncio
async def test_api_import_params_and_pagination(session):
    import httpx, json
    from fastapi import FastAPI
    from app.routers.collector import router
    from app.auth import get_current_user
    from app.database import get_db
    app = FastAPI(); app.include_router(router)
    async def db(): yield session
    app.dependency_overrides[get_db] = db
    app.dependency_overrides[get_current_user] = lambda: SimpleNamespace(id=1, role='admin')
    params={'sources':['import'], 'rights_confirmed':True, 'rights_basis':'Fixture', 'crawl_websites':False}
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url='http://test') as client:
        response = await client.post('/api/collector/import', data={'params':json.dumps(params)}, files={'file':('fixture.csv',b'name,website,email\nFixture,https://fixture.ru,office@fixture.ru\n','text/csv')})
        assert response.status_code == 200, response.text
        assert response.json()['rows'] == 1
        for index in range(51): await save_company(session, record(index), source='import', rights=RIGHTS)
        response = await client.get('/api/collector/companies?page=2&has_email=true')
        assert response.status_code == 200 and len(response.json()['items']) == 1
        assert response.json()['total'] == 51
